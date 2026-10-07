#!/usr/bin/env python3
# Gera a planilha de indices de desempenho (T2) a partir das duas simulacoes.
#
# Uso: python3 gerar_planilha.py
# Saida: planilhas/indices_oficina.xlsx e planilhas/indices_oficina.csv

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import simulador as sim

BASE = os.path.dirname(os.path.abspath(__file__))
ATUAL = os.path.join(BASE, 'modelos', 'oficina.yml')
MELHOR = os.path.join(BASE, 'modelos', 'oficina_melhorada.yml')

AZUL = '1F4E79'
CINZA = 'D9E1F2'
VERDE = 'C6EFCE'
VERMELHO = 'FFC7CE'
AMARELO = 'FFEB9C'

BORDA = Border(*[Side(style='thin', color='BFBFBF')] * 4)


def coleta(caminho):
    esc, n = sim.simula(caminho)
    tg = esc.tempo_global()
    linhas = []
    for fila in esc.filas.values():
        idx = sim.indices(fila, tg)
        linhas.append({
            'fila': fila.nome,
            'notacao': sim.notacao(fila),
            'servidores': fila.servidores,
            'capacidade': fila.capacidade,
            'chegadas_externas': fila.chegadas_externas,
            'atendidos': fila.atendidos,
            'perdas': fila.perdidos,
            'saidas': fila.saidas,
            'es': idx['es'],
            'mu': idx['mu'],
            'n': idx['n'],
            'd': idx['d'],
            'u': idx['u'],
            'w': idx['w'],
            'estados': dict(fila.estatistica),
        })
    return {'tg': tg, 'n': n, 'filas': linhas}


def estilo_cabecalho(celula):
    celula.font = Font(bold=True, color='FFFFFF')
    celula.fill = PatternFill('solid', fgColor=AZUL)
    celula.alignment = Alignment(horizontal='center', vertical='center')
    celula.border = BORDA


def ajusta_largura(ws):
    for coluna in ws.columns:
        maior = 0
        letra = get_column_letter(coluna[0].column)
        for celula in coluna:
            if celula.value is not None:
                maior = max(maior, len(str(celula.value)))
        ws.column_dimensions[letra].width = min(maior + 3, 34)


def main():
    atual = coleta(ATUAL)
    melhor = coleta(MELHOR)
    wb = Workbook()

    # ------------------------------------------------------------------ Resumo
    ws = wb.active
    ws.title = 'Comparativo'
    ws['A1'] = 'T2 | Avaliacao de um Modelo Simulado - Oficina Mecanica'
    ws['A1'].font = Font(bold=True, size=14, color=AZUL)
    ws['A2'] = 'Indices de desempenho por fila (10 sementes x 100.000 aleatorios)'
    ws['A2'].font = Font(italic=True, color='595959')

    cabecalho = ['Fila', 'Cenario', 'Notacao', 'Servidores', 'Capacidade',
                 'Chegadas externas', 'Atendidos', 'Perdas', 'Saidas',
                 'E[S] (min)', 'MU (cli/min)', 'Populacao N', 'Vazao D (cli/min)',
                 'Utilizacao U', 'Tempo resposta W (min)']
    linha = 4
    for c, titulo in enumerate(cabecalho, start=1):
        celula = ws.cell(row=linha, column=c, value=titulo)
        estilo_cabecalho(celula)
    linha += 1
    for cenario, dados, cor in (('Atual', atual, VERMELHO),
                                ('Melhorada', melhor, VERDE)):
        for f in dados['filas']:
            valores = [f['fila'], cenario, f['notacao'], f['servidores'],
                       f['capacidade'], f['chegadas_externas'], f['atendidos'],
                       f['perdas'], f['saidas'], round(f['es'], 4),
                       round(f['mu'], 6), round(f['n'], 4), round(f['d'], 6),
                       round(f['u'], 4), round(f['w'], 4)]
            for c, valor in enumerate(valores, start=1):
                celula = ws.cell(row=linha, column=c, value=valor)
                celula.border = BORDA
                if c in (14,):
                    celula.number_format = '0.00%'
                if cenario == 'Atual' and c in (8, 15):
                    celula.fill = PatternFill('solid', fgColor=AMARELO)
                if cenario == 'Melhorada' and c in (8, 15):
                    celula.fill = PatternFill('solid', fgColor=cor)
            linha += 1

    # Linha de variacao (melhorada vs atual) por fila
    linha += 1
    ws.cell(row=linha, column=1, value='Variacao (Melhorada - Atual)').font = Font(bold=True)
    linha += 1
    for c, titulo in enumerate(['Fila', 'd Perdas', 'd N', 'd D', 'd U', 'd W'], start=1):
        estilo_cabecalho(ws.cell(row=linha, column=c, value=titulo))
    linha += 1
    mapa_atual = {f['fila']: f for f in atual['filas']}
    for f in melhor['filas']:
        a = mapa_atual[f['fila']]
        valores = [f['fila'], f['perdas'] - a['perdas'], round(f['n'] - a['n'], 4),
                   round(f['d'] - a['d'], 6), round(f['u'] - a['u'], 4),
                   round(f['w'] - a['w'], 4)]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=c, value=valor)
            celula.border = BORDA
        linha += 1
    ajusta_largura(ws)

    # ------------------------------------------------- distribuicao de estados
    for titulo, dados in (('Estados atual', atual), ('Estados melhorada', melhor)):
        ws = wb.create_sheet(titulo)
        ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, color=AZUL)
        linha = 3
        for f in dados['filas']:
            ws.cell(row=linha, column=1, value='%s (%s)' % (f['fila'], f['notacao']))
            ws.cell(row=linha, column=1).font = Font(bold=True)
            linha += 1
            for c, t in enumerate(['Estado', 'Tempo acumulado (min)', 'Probabilidade'], start=1):
                estilo_cabecalho(ws.cell(row=linha, column=c, value=t))
            linha += 1
            for estado in sorted(f['estados']):
                tempo = f['estados'][estado]
                celula = ws.cell(row=linha, column=2, value=round(tempo, 4))
                celula.border = BORDA
                celula = ws.cell(row=linha, column=3, value=tempo / dados['tg'])
                celula.number_format = '0.00%'
                celula.border = BORDA
                ws.cell(row=linha, column=1, value=estado).border = BORDA
                linha += 1
            linha += 1
        ajusta_largura(ws)

    # -------------------------------------------------------------------- csv
    caminho_xlsx = os.path.join(BASE, 'planilhas', 'indices_oficina.xlsx')
    wb.save(caminho_xlsx)

    caminho_csv = os.path.join(BASE, 'planilhas', 'indices_oficina.csv')
    with open(caminho_csv, 'w', newline='', encoding='utf-8') as arquivo:
        escritor = csv.writer(arquivo, delimiter=';')
        escritor.writerow(cabecalho)
        for cenario, dados in (('Atual', atual), ('Melhorada', melhor)):
            for f in dados['filas']:
                escritor.writerow([f['fila'], cenario, f['notacao'], f['servidores'],
                                   f['capacidade'], f['chegadas_externas'], f['atendidos'],
                                   f['perdas'], f['saidas'], round(f['es'], 4),
                                   round(f['mu'], 6), round(f['n'], 4), round(f['d'], 6),
                                   round(f['u'], 4), round(f['w'], 4)])
    print('gerado:', caminho_xlsx)
    print('gerado:', caminho_csv)

    # ---------------------------------------------------- resumo no terminal
    for rotulo, dados in (('ATUAL', atual), ('MELHORADA', melhor)):
        print('\n== %s ==  TG=%.2f min' % (rotulo, dados['tg']))
        print('%-12s %6s %8s %8s %8s %8s' % ('fila', 'perdas', 'N', 'D', 'U', 'W'))
        for f in dados['filas']:
            print('%-12s %6d %8.3f %8.4f %7.2f%% %8.2f' % (
                f['fila'], f['perdas'], f['n'], f['d'], 100 * f['u'], f['w']))


if __name__ == '__main__':
    main()
