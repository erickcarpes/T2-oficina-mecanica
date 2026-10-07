#!/usr/bin/env python3
# Gera a apresentacao do T2 (PPTX) a partir das simulacoes e dos desenhos.
#
# Uso: python3 gerar_apresentacao.py
# Saida: apresentacao/T2_Oficina_Mecanica.pptx

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

import simulador as sim

BASE = os.path.dirname(os.path.abspath(__file__))
APRES = os.path.join(BASE, 'apresentacao')
ATUAL = os.path.join(BASE, 'modelos', 'oficina.yml')
MELHOR = os.path.join(BASE, 'modelos', 'oficina_melhorada.yml')

AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_CLARO = RGBColor(0xDE, 0xEB, 0xF7)
CINZA = RGBColor(0x59, 0x59, 0x59)
VERMELHO = RGBColor(0xC0, 0x00, 0x00)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def br(valor, casas=2):
    return ('%.*f' % (casas, valor)).replace('.', ',')


def coleta(caminho):
    esc, n = sim.simula(caminho)
    tg = esc.tempo_global()
    filas = {}
    for fila in esc.filas.values():
        idx = sim.indices(fila, tg)
        idx['perdas'] = fila.perdidos
        idx['estados'] = dict(fila.estatistica)
        idx['notacao'] = sim.notacao(fila)
        idx['servidores'] = fila.servidores
        idx['capacidade'] = fila.capacidade
        filas[fila.nome] = idx
    return {'tg': tg, 'filas': filas}


def caixa_texto(slide, x, y, w, h, titulo, linhas, tam_titulo=26, tam_texto=16):
    if titulo:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = titulo
        run.font.size = Pt(tam_titulo)
        run.font.bold = True
        run.font.color.rgb = AZUL
        y += 0.95
    if linhas:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, linha in enumerate(linhas):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            if isinstance(linha, tuple):
                texto, nivel, cor = linha
            else:
                texto, nivel, cor = linha, 0, CINZA
            p.level = nivel
            run = p.add_run()
            run.text = texto
            run.font.size = Pt(tam_texto - 2 * nivel)
            run.font.color.rgb = cor
    return tb


def slide_branco(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rodape(slide, numero, total):
    tb = slide.shapes.add_textbox(Inches(9.0), Inches(7.0), Inches(3.5), Inches(0.35))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = 'T2 | Oficina Mecanica  ·  %d/%d' % (numero, total)
    run.font.size = Pt(10)
    run.font.color.rgb = CINZA


def tabela(slide, x, y, w, h, dados, larguras=None, tam=12, cabecalho=True):
    linhas = len(dados)
    colunas = len(dados[0])
    shape = slide.shapes.add_table(linhas, colunas, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    tabela_ = shape.table
    if larguras:
        for i, larg in enumerate(larguras):
            tabela_.columns[i].width = Inches(larg)
    for r, linha in enumerate(dados):
        for c, valor in enumerate(linha):
            celula = tabela_.cell(r, c)
            celula.text = str(valor)
            celula.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in celula.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.size = Pt(tam)
                    if r == 0 and cabecalho:
                        run.font.bold = True
                        run.font.color.rgb = BRANCO
                    else:
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            if r == 0 and cabecalho:
                celula.fill.solid()
                celula.fill.fore_color.rgb = AZUL
            elif r % 2 == 0:
                celula.fill.solid()
                celula.fill.fore_color.rgb = AZUL_CLARO
    return tabela_


def main():
    atual = coleta(ATUAL)
    melhor = coleta(MELHOR)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 10
    n = 0

    # ---------------------------------------------------------------- Capa
    n += 1
    slide = slide_branco(prs)
    faixa = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(3.2))
    faixa.fill.solid()
    faixa.fill.fore_color.rgb = AZUL
    faixa.line.fill.background()
    tf = faixa.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.8)
    tf.margin_top = Inches(0.8)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = 'T2 | Avaliacao de Aprendizagem'
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0xBD, 0xD7, 0xEE)
    p2 = tf.add_paragraph()
    run = p2.add_run()
    run.text = 'Avaliacao de um Modelo Simulado'
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = BRANCO
    p3 = tf.add_paragraph()
    run = p3.add_run()
    run.text = 'Oficina Mecanica: analise e melhoria de uma rede de filas com realimentacao'
    run.font.size = Pt(20)
    run.font.color.rgb = BRANCO

    caixa_texto(slide, 0.8, 3.7, 8.0, 3.0, '',
                [('Disciplina: Simulacao e Metodos Analiticos', 0, CINZA),
                 ('Unidade de Aprendizagem: UA2 | Fundamentos de Simulacao por Computador', 0, CINZA),
                 ('Modulo de Aprendizagem: M9 | Avaliacao de um Modelo Simulado', 0, CINZA),
                 ('Grupo: T2 - 690 - 16', 0, CINZA)] + [('', 0, CINZA)] +
                [('Integrantes:', 0, AZUL),
                 ('Integrante 1  ·  Integrante 2  ·  Integrante 3  ·  Integrante 4', 0, CINZA),
                 ('Integrante 5  ·  Integrante 6', 0, CINZA)] +
                [('', 0, CINZA), ('Porto Alegre, outubro de 2026', 0, CINZA)],
                tam_texto=16)
    rodape(slide, n, total)

    # ---------------------------------------------------- Descricao da realidade
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.4, 12.0, 5.9, 'Descricao da realidade a ser simulada',
                [('Uma oficina mecanica recebe veiculos para reparo. O cliente chega, abre uma ordem de servico '
                  'na RECEPCAO (orcamento e cadastro) e o carro segue para a OFICINA, onde e reparado. Ao terminar, '
                  'passa pelo CONTROLE DE QUALIDADE (teste do servico).', 0, CINZA),
                 ('', 0, CINZA),
                 ('A rede tem realimentacao (nao e um fluxo em tandem):', 0, AZUL),
                 ('Retrabalho: se o teste reprova, o veiculo VOLTA da QUALIDADE para a OFICINA.', 1, CINZA),
                 ('Pecas/orcamento: parte dos servicos exige orcamento extra e VOLTA da OFICINA para a RECEPCAO.', 1, CINZA),
                 ('', 0, CINZA),
                 ('Hipotese de modelagem:', 0, AZUL),
                 ('Chegadas do exterior apenas na RECEPCAO; intervalos e tempos de servico uniformes U(a,b).', 1, CINZA),
                 ('Disciplina FCFS; capacidade finita em todas as filas (chegada com fila cheia = perda).', 1, CINZA),
                 ('Simulacao: 10 sementes x 100.000 numeros pseudoaleatorios (gerador congruente linear).', 1, CINZA)],
                tam_texto=16)
    rodape(slide, n, total)

    # ------------------------------------------------------------ Modelo atual
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.25, 12.0, 0.7, 'Modelo (desenho) da rede de filas - situacao ATUAL',
                [], tam_titulo=26)
    slide.shapes.add_picture(os.path.join(APRES, 'rede_filas_atual.png'),
                             Inches(0.55), Inches(1.15), width=Inches(12.2))
    rodape(slide, n, total)

    # ------------------------------------------------------- Parametros/Kendall
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.3, 12.0, 0.7, 'Descricao das filas (notacao de Kendall) e roteamento',
                [], tam_titulo=24)
    dados = [
        ['Fila', 'Notacao', 'Chegadas ext.', 'Atendimento', 'Servidores', 'Capacidade'],
        ['RECEPCAO', 'G/G/1/8', 'U(6;12), 1a em 5,0', 'U(4;8)', '1', '8'],
        ['OFICINA', 'G/G/3/12', '--', 'U(20;35)', '3', '12'],
        ['QUALIDADE', 'G/G/1/6', '--', 'U(4;8)', '1', '6'],
    ]
    tabela(slide, 0.7, 1.4, 12.0, 1.8, dados,
           larguras=[2.2, 1.8, 2.4, 2.0, 1.8, 1.8], tam=14)
    caixa_texto(slide, 0.7, 3.5, 12.0, 3.4, 'Probabilidades de roteamento',
                [('RECEPCAO -> OFICINA: 95%      RECEPCAO -> SAIDA: 5%', 0, CINZA),
                 ('OFICINA -> QUALIDADE: 90%    OFICINA -> RECEPCAO: 5% (pecas)    OFICINA -> SAIDA: 5%', 0, CINZA),
                 ('QUALIDADE -> OFICINA: 15% (retrabalho)     QUALIDADE -> SAIDA: 85%', 0, CINZA),
                 ('', 0, CINZA),
                 ('A soma das probabilidades de cada fila e 1,0 (incluindo a saida da rede).', 0, AZUL)],
                tam_texto=15)
    rodape(slide, n, total)

    # -------------------------------------------------- Resultados - atual
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.3, 12.0, 0.7, 'Resultados da simulacao - modelo ATUAL', [], tam_titulo=24)
    a = atual['filas']
    dados = [
        ['Fila', 'Perdas', 'N (populacao)', 'D (vazao/min)', 'U (utilizacao)', 'W (min)'],
        ['RECEPCAO', str(a['RECEPCAO']['perdas']), br(a['RECEPCAO']['n']), br(a['RECEPCAO']['d'], 4),
         br(100 * a['RECEPCAO']['u']) + '%', br(a['RECEPCAO']['w'])],
        ['OFICINA', str(a['OFICINA']['perdas']), br(a['OFICINA']['n']), br(a['OFICINA']['d'], 4),
         br(100 * a['OFICINA']['u']) + '%', br(a['OFICINA']['w'])],
        ['QUALIDADE', str(a['QUALIDADE']['perdas']), br(a['QUALIDADE']['n']), br(a['QUALIDADE']['d'], 4),
         br(100 * a['QUALIDADE']['u']) + '%', br(a['QUALIDADE']['w'])],
    ]
    tabela(slide, 0.7, 1.5, 12.0, 1.9, dados,
           larguras=[2.2, 1.4, 2.1, 2.2, 2.1, 2.0], tam=14)
    caixa_texto(slide, 0.7, 3.7, 12.0, 3.3, '',
                [('Distribuicao de probabilidades dos estados (OFICINA):', 0, AZUL),
                 ('P(0)=0,01%  P(5)=0,08%  P(9)=7,72%  P(10)=24,69%  P(11)=39,21%  P(12)=25,22%', 0, CINZA),
                 ('A fila passa a maior parte do tempo lotada (estados 10 a 12 somam ~89%).', 0, CINZA),
                 ('', 0, CINZA),
                 ('Outras metricas:', 0, AZUL),
                 ('Tempo global da simulacao (10 sementes): ' + br(atual['tg']) + ' min.', 0, CINZA),
                 ('Todas as filas acumulam tempo de estado a cada evento; P(estado) = tempo no estado / tempo global.', 0, CINZA)],
                tam_texto=15)
    rodape(slide, n, total)

    # ---------------------------------------------------------- Gargalo
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.3, 12.0, 0.7, 'Identificacao do gargalo', [], tam_titulo=24)
    caixa_texto(slide, 0.7, 1.4, 12.0, 5.4, '',
                [('A OFICINA esta saturada: U = 99,98% (todos os 3 mecanicos ocupados todo o tempo).', 0, VERMELHO),
                 ('Populacao media N = ' + br(a['OFICINA']['n']) + ' de 12 vagas; tempo de resposta W = '
                  + br(a['OFICINA']['w']) + ' min (~2h por veiculo, quase so espera).', 0, VERMELHO),
                 ('Foram perdidos ' + str(a['OFICINA']['perdas']) + ' clientes por fila cheia na OFICINA.', 0, VERMELHO),
                 ('', 0, CINZA),
                 ('A carga oferecida supera a capacidade instalada:', 0, AZUL),
                 ('Taxa de atendimento: MU = 1/E[S] = 1/27,5 = 0,03636 cliente/min por mecanico.', 1, CINZA),
                 ('Capacidade da oficina: 3 x 0,03636 = 0,1091 cliente/min.', 1, CINZA),
                 ('Chegada efetiva a oficina: ~0,1291 cliente/min  ->  rho = 0,1291/0,1091 = 1,18 (> 1).', 1, CINZA),
                 ('', 0, CINZA),
                 ('Conclusao: e preciso aumentar a capacidade da OFICINA e/ou reduzir a carga que chega ate ela.', 0, VERDE)],
                tam_texto=16)
    rodape(slide, n, total)

    # -------------------------------------------------------- Proposta melhoria
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.25, 12.0, 0.7, 'Proposta de melhoria e modelo MELHORADO',
                [], tam_titulo=24)
    caixa_texto(slide, 0.7, 1.05, 12.0, 1.0, '',
                [('Alteracoes no modelo (destacadas em vermelho/verde no desenho):', 0, AZUL),
                 ('OFICINA: 3 -> 4 mecanicos, capacidade 12 -> 14, atendimento U(20;35) -> U(18;30)  '
                  '(contratacao + elevador/ferramenta).', 1, VERMELHO),
                 ('QUALIDADE: capacidade 6 -> 8 e retrabalho 15% -> 10% (treinamento reduz devolucoes).', 1, VERDE),
                 ('RECEPCAO: capacidade 8 -> 10 (evita perdas de ordens na chegada).', 1, CINZA)],
                tam_texto=15)
    slide.shapes.add_picture(os.path.join(APRES, 'rede_filas_melhorada.png'),
                             Inches(0.55), Inches(2.35), width=Inches(12.2))
    rodape(slide, n, total)

    # --------------------------------------------------------- Comparacao
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.3, 12.0, 0.7, 'Comparacao dos resultados (antes x depois)', [], tam_titulo=24)
    m = melhor['filas']
    dados = [
        ['Fila', 'Metrica', 'Atual', 'Melhorada', 'Variacao'],
        ['OFICINA', 'Perdas', str(a['OFICINA']['perdas']), str(m['OFICINA']['perdas']), 'zerou'],
        ['OFICINA', 'N (populacao)', br(a['OFICINA']['n']), br(m['OFICINA']['n']),
         '-' + br(a['OFICINA']['n'] - m['OFICINA']['n'])],
        ['OFICINA', 'D (vazao/min)', br(a['OFICINA']['d'], 4), br(m['OFICINA']['d'], 4),
         '+' + br(100 * (m['OFICINA']['d'] / a['OFICINA']['d'] - 1), 1) + '%'],
        ['OFICINA', 'U (utilizacao)', br(100 * a['OFICINA']['u']) + '%', br(100 * m['OFICINA']['u']) + '%',
         '-' + br(100 * (a['OFICINA']['u'] - m['OFICINA']['u']), 1) + ' p.p.'],
        ['OFICINA', 'W (min)', br(a['OFICINA']['w']), br(m['OFICINA']['w']),
         '-' + br(100 * (1 - m['OFICINA']['w'] / a['OFICINA']['w']), 1) + '%'],
    ]
    tabela(slide, 0.5, 1.35, 6.6, 4.6, dados,
           larguras=[1.4, 1.5, 1.2, 1.3, 1.2], tam=11)

    chart_data = CategoryChartData()
    chart_data.categories = ['RECEPCAO', 'OFICINA', 'QUALIDADE']
    chart_data.add_series('W atual (min)', (a['RECEPCAO']['w'], a['OFICINA']['w'], a['QUALIDADE']['w']))
    chart_data.add_series('W melhorada (min)', (m['RECEPCAO']['w'], m['OFICINA']['w'], m['QUALIDADE']['w']))
    grafico = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.3), Inches(1.35),
                                     Inches(5.6), Inches(4.6), chart_data).chart
    grafico.has_title = True
    grafico.chart_title.text_frame.text = 'Tempo de resposta W por fila (min)'
    grafico.has_legend = True
    grafico.legend.position = XL_LEGEND_POSITION.BOTTOM
    grafico.legend.include_in_layout = False
    rodape(slide, n, total)

    # ---------------------------------------------------------- Conclusoes
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.35, 12.0, 0.8, 'Conclusoes finais', [], tam_titulo=26)
    caixa_texto(slide, 0.7, 1.5, 12.0, 5.4, '',
                [('A simulacao revelou que a OFICINA e o gargalo: saturada (U~100%), com longas esperas e 21.406 perdas.', 0, CINZA),
                 ('A melhoria (mais um mecanico, processo mais rapido e menor retrabalho) zerou as perdas da oficina.', 0, CINZA),
                 ('O tempo de resposta na OFICINA caiu de ' + br(a['OFICINA']['w']) + ' min para ' + br(m['OFICINA']['w'])
                  + ' min (' + br(100 * (1 - m['OFICINA']['w'] / a['OFICINA']['w']), 1) + '% melhor).', 0, VERDE),
                 ('A vazao da oficina aumentou (mais veiculos entregues por minuto) com utilizacao confortavel (73%).', 0, CINZA),
                 ('A QUALIDADE absorveu o maior fluxo sem perdas; e a RECEPCAO praticamente nao mudou.', 0, CINZA),
                 ('', 0, CINZA),
                 ('Trabalhos futuros:', 0, AZUL),
                 ('Analisar custo x beneficio (custo do mecanico extra e do equipamento vs. perda de clientes).', 1, CINZA),
                 ('Executar varias replicacoes e intervalos de confianca para os indices.', 1, CINZA)],
                tam_texto=16)
    rodape(slide, n, total)

    # ---------------------------------------------------- Arquivos e reproducao
    n += 1
    slide = slide_branco(prs)
    caixa_texto(slide, 0.7, 0.35, 12.0, 0.8, 'Arquivos e reproducao', [], tam_titulo=26)
    caixa_texto(slide, 0.7, 1.4, 12.0, 5.5, '',
                [('O pacote contem:', 0, AZUL),
                 ('modelos/oficina.yml e modelos/oficina_melhorada.yml - modelos antes e depois da melhoria.', 1, CINZA),
                 ('simulador.py - simulador de rede de filas (nosso, do T1) com o calculo dos indices.', 1, CINZA),
                 ('resultados/oficina_atual.txt e oficina_melhorada.txt - saidas completas (tempos e probabilidades por estado).', 1, CINZA),
                 ('planilhas/indices_oficina.xlsx - calculo dos indices (populacao, vazao, utilizacao, tempo de resposta).', 1, CINZA),
                 ('apresentacao/ - este PDF/PPTX e os desenhos da rede.', 1, CINZA),
                 ('', 0, CINZA),
                 ('Como reproduzir:', 0, AZUL),
                 ('pip install pyyaml', 1, CINZA),
                 ('python3 simulador.py modelos/oficina.yml', 1, CINZA),
                 ('python3 simulador.py modelos/oficina_melhorada.yml', 1, CINZA),
                 ('Formulas: MU = 1/E[S];  N = soma k.P(k);  D = MU.Ns;  U = Ns/servidores;  W = N/D.', 0, AZUL)],
                tam_texto=15)
    rodape(slide, n, total)

    caminho = os.path.join(APRES, 'T2_Oficina_Mecanica.pptx')
    prs.save(caminho)
    print('gerado:', caminho)
    print('slides:', len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == '__main__':
    main()
